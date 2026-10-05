import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def build_word_doc(md_filepath, output_filepaths):
    doc = docx.Document()
    
    # Page Margins (1 inch everywhere)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    NAVY_COLOR = RGBColor(15, 27, 42)      # #0f1b2a
    BLUE_COLOR = RGBColor(9, 114, 211)     # #0972d3
    DARK_GRAY  = RGBColor(22, 25, 31)      # #16191f
    SLATE_GRAY = RGBColor(84, 91, 100)     # #545b64

    # Add Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("⚡ Intelligent EV Charging & Resource Management System")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY_COLOR
    run_title.font.name = "Arial"

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Executive Technical Documentation & Operational Staff Manual")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = BLUE_COLOR
    run_sub.font.name = "Arial"

    doc.add_paragraph()  # spacing

    with open(md_filepath, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table(rows):
        if not rows:
            return
        parsed_table = []
        for r in rows:
            cols = [c.strip() for c in r.strip('|').split('|')]
            parsed_table.append(cols)

        if len(parsed_table) < 2:
            return

        header = parsed_table[0]
        data = [r for r in parsed_table[1:] if not (len(r) > 0 and '---' in r[0])]

        table = doc.add_table(rows=len(data) + 1, cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header formatting
        hdr_cells = table.rows[0].cells
        for idx, title_text in enumerate(header):
            hdr_cells[idx].text = title_text.replace('**', '').replace('$', '')
            set_cell_background(hdr_cells[idx], "0F1B2A")
            for p in hdr_cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                    r.font.size = Pt(10)
                    r.font.name = "Arial"

        # Data rows formatting
        for row_idx, row_data in enumerate(data):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                if col_idx < len(row_cells):
                    row_cells[col_idx].text = cell_value.replace('**', '').replace('*', '')
                    set_cell_background(row_cells[col_idx], bg_color)
                    for p in row_cells[col_idx].paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = DARK_GRAY
                            r.font.name = "Arial"

        doc.add_paragraph()  # spacing after table

    for line in lines:
        raw = line.strip()

        # Handle Code blocks
        if raw.startswith("```"):
            if in_code_block:
                in_code_block = False
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.4)
                p_code.paragraph_format.right_indent = Inches(0.4)
                run_c = p_code.add_run("\n".join(code_lines))
                run_c.font.name = "Consolas"
                run_c.font.size = Pt(9)
                run_c.font.color.rgb = DARK_GRAY
                code_lines = []
            else:
                if in_table:
                    flush_table(table_rows)
                    table_rows = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line.rstrip("\n"))
            continue

        # Handle Tables
        if raw.startswith("|") and raw.endswith("|"):
            in_table = True
            table_rows.append(raw)
            continue
        elif in_table:
            flush_table(table_rows)
            table_rows = []
            in_table = False

        if not raw:
            continue

        # Headings
        if raw.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(raw[2:].replace('**', ''))
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = NAVY_COLOR
            run.font.name = "Arial"
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)

        elif raw.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(raw[3:].replace('**', ''))
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = BLUE_COLOR
            run.font.name = "Arial"
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)

        elif raw.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(raw[4:].replace('**', ''))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = NAVY_COLOR
            run.font.name = "Arial"
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)

        elif raw.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(raw[2:].replace('**', ''))
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = SLATE_GRAY
            run.font.name = "Arial"

        elif raw.startswith("- ") or raw.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(raw[2:].replace('**', ''))
            run.font.size = Pt(10.5)
            run.font.color.rgb = DARK_GRAY
            run.font.name = "Arial"

        elif raw.startswith("1. ") or raw.startswith("2. ") or raw.startswith("3. ") or raw.startswith("4. ") or raw.startswith("5. ") or raw.startswith("6. "):
            p = doc.add_paragraph(style='List Number')
            run = p.add_run(raw[3:].replace('**', ''))
            run.font.size = Pt(10.5)
            run.font.color.rgb = DARK_GRAY
            run.font.name = "Arial"

        elif raw.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)

        else:
            p = doc.add_paragraph()
            run = p.add_run(raw.replace('**', ''))
            run.font.size = Pt(10.5)
            run.font.color.rgb = DARK_GRAY
            run.font.name = "Arial"
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table(table_rows)

    for out_path in output_filepaths:
        doc.save(out_path)
        print(f"✅ Saved Word Document to: {out_path}")

if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    workspace_root = os.path.abspath(os.path.join(current_dir, ".."))
    
    md_file = os.path.join(current_dir, "DOCUMENTATION.md")
    if not os.path.exists(md_file):
        md_file = os.path.join(workspace_root, "DOCUMENTATION.md")

    out_files = [
        os.path.join(current_dir, "Documentation.docx"),
        os.path.join(workspace_root, "Intelligent_EV_Charging_System_Documentation.docx")
    ]
    print(f"Reading documentation from: {md_file}")
    build_word_doc(md_file, out_files)

